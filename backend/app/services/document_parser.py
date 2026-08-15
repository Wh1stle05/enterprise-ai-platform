import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader


class DocumentParseError(ValueError):
    pass


class UnsupportedDocumentError(DocumentParseError):
    pass


class EmptyDocumentError(DocumentParseError):
    pass


@dataclass(frozen=True)
class SourceSection:
    text: str
    locator: str


def _normalize(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n")).strip()


def _parse_pdf(path: Path) -> list[SourceSection]:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocumentParseError("PDF is encrypted")
        return [
            SourceSection(_normalize(page.extract_text() or ""), f"page {i}")
            for i, page in enumerate(reader.pages, 1)
        ]
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("Unable to parse PDF") from exc


def _parse_docx(path: Path) -> list[SourceSection]:
    try:
        doc = WordDocument(str(path))
        return [
            SourceSection(_normalize(p.text), f"paragraph {i}")
            for i, p in enumerate(doc.paragraphs, 1)
        ]
    except Exception as exc:
        raise DocumentParseError("Unable to parse Word document") from exc


def _parse_xlsx(path: Path) -> list[SourceSection]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections = []
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [str(value) for value in row if value is not None]
                sections.append(
                    SourceSection(
                        _normalize(" | ".join(values)),
                        f"sheet {sheet.title} row {row_number}",
                    )
                )
        return sections
    except Exception as exc:
        raise DocumentParseError("Unable to parse Excel workbook") from exc


def _parse_markdown(path: Path) -> list[SourceSection]:
    try:
        return [SourceSection(_normalize(path.read_text(encoding="utf-8")), "document")]
    except UnicodeDecodeError as exc:
        raise DocumentParseError("Markdown is not valid UTF-8") from exc
    except OSError as exc:
        raise DocumentParseError("Unable to read Markdown document") from exc


def parse_document(path: Path) -> list[SourceSection]:
    parsers = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx,
        ".md": _parse_markdown,
    }
    parser = parsers.get(path.suffix.lower())
    if parser is None:
        raise UnsupportedDocumentError(path.suffix.lower())
    sections = [section for section in parser(path) if section.text.strip()]
    if not sections:
        raise EmptyDocumentError("Document contains no extractable text")
    return sections
